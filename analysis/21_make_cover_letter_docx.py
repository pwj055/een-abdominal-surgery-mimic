# -*- coding: utf-8 -*-
"""Build the Critical Care cover letter as .docx.

Output: ../cover_letter_criticalcare.docx
"""
import os
import re

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.abspath(os.path.join(HERE, "..", "cover_letter_criticalcare.docx"))

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.0
_rpr = normal.element.get_or_add_rPr()
_rfonts = _rpr.get_or_add_rFonts()
_rfonts.set(qn("w:eastAsia"), "SimSun")


def add_rich(p, text):
    for part in re.split(r"(\*\*.*?\*\*|\*[^*]+?\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1])
            r.italic = True
        else:
            p.add_run(part)
    return p


def para(text="", align=None, space_after=6, size=None):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    add_rich(p, text)
    if size:
        for r in p.runs:
            r.font.size = Pt(size)
    return p


def heading(text):
    p = para(text, space_after=4)
    for r in p.runs:
        r.bold = True
    return p


# Letter body
para("October 18, 2026", space_after=12)
para("Dear Editor,", space_after=12)

heading("Re: Submission of a research manuscript to Critical Care")

para("**Title:** Early versus Late Enteral Nutrition and Infectious Complications after Major "
     "Abdominal Surgery in the Intensive Care Unit: A Target Trial Emulation Using MIMIC-IV")
para("**Article type:** Research")
para("**Corresponding author:** Yehong Han, jianghao-163@163.com", space_after=12)

para("We are pleased to submit the above manuscript for consideration for publication in "
     "*Critical Care*.", align="justify")

heading("Why this manuscript is suited to Critical Care")
para("The optimal timing of enteral nutrition after major abdominal surgery in the ICU remains "
     "an unsettled, everyday decision in critical care practice, yet the existing evidence base "
     "is dominated by trials conducted in medical and mixed ICUs that largely excluded surgical "
     "patients. We address this gap using the target-trial-emulation framework applied to the "
     "MIMIC-IV critical care database, with two prespecified co-primary comparisons that "
     "explicitly separate the *timing* estimand (early vs. late feeding among patients ultimately "
     "fed) from the *policy* estimand (early feeding vs. usual care in the full cohort). The key "
     "finding — a lower risk of antibiotic-treated infection with earlier feeding in the "
     "eventually-fed stratum, with a consistent direction but imprecise magnitude — is directly "
     "relevant to intensivists, surgeons, and nutrition support teams, and is reported "
     "transparently together with a full causal-inference sensitivity program (overlap weighting, "
     "IPTW, AIPW, propensity-score matching, negative-control outcomes, E-values, and same-center "
     "temporal and multicenter triangulations).", align="justify")
para("We believe the methodological approach — separating principal-stratum and policy estimands "
     "to avoid a spurious reversal of effect direction — will be of broad interest to the "
     "readership of *Critical Care*, and it provides a template for observational feeding-timing "
     "research beyond this dataset.", align="justify")

heading("Key declarations")
para("**Originality and exclusivity.** The manuscript is original, has not been published "
     "previously, is not under consideration by another journal, and will not be submitted "
     "elsewhere while under review at *Critical Care*. All authors have read and approved the "
     "manuscript, agree to its submission, and meet the ICMJE criteria for authorship.", align="justify")
para("**Reporting standards.** The study is reported in accordance with the STROBE statement and "
     "its RECORD extension; completed checklists are provided as Additional Files 1 and 2. The "
     "study protocol and analysis plan were time-stamped and registered at the Open Science "
     "Framework (https://osf.io/g39vw; DOI 10.17605/OSF.IO/G39VW) before outcome modeling.", align="justify")
para("**Ethics.** The study used publicly available, de-identified data (MIMIC-IV, MIMIC-III, "
     "eICU-CRD) under credentialed PhysioNet data-use agreements; informed consent was waived by "
     "the respective institutional review boards.", align="justify")
para("**Competing interests.** The authors declare that they have no competing interests.", align="justify")
para("**Funding.** This work was supported by the Zhejiang Provincial Health Industry Science and Technology Program Provincial–Ministerial Jointly Funded Project (grant number WKJ-ZJ-26079).", align="justify")
para("**Data availability.** The clinical databases are available from PhysioNet subject to "
     "credentialed data-use agreements; analysis code is public on GitHub and archived on Zenodo "
     "(https://doi.org/10.5281/zenodo.23446232).", align="justify")

heading("ORCID")
para("Peiwu Jiang: 0000-0002-5262-3424", align="justify")

para("Thank you for considering our manuscript. We look forward to your response.", align="justify")

para("Yours sincerely,", space_after=4)
para("Yehong Han (on behalf of all authors)")
para("Corresponding author")
para("Department of General Surgery, Hangzhou TCM Hospital Affiliated to Zhejiang Chinese Medical University, Hangzhou, Zhejiang, China.")
para("jianghao-163@163.com")

doc.save(OUT)
print(f"Saved {OUT}")
